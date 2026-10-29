# flake8: noqa
# pyright: reportUndefinedVariable=false
# =====================================================================================
# COMPRAS · MEJORES PRECIOS (análisis de compras)
# A qué precio compramos cada producto, insumo o servicio y a quién: mejor precio, último, promedio,
# ranking de proveedores, ahorro potencial, estadísticas, gráficos y Excel.
# Este archivo NO se importa: app.py lo carga en su lugar (lista MODULOS_EN_ARCHIVOS), como si
# su código estuviera escrito ahí. Usa directamente los modelos, variables y funciones de app.py
# y de los archivos que se cargan antes. Ver MAPA.md.
# =====================================================================================
# Fuente: ítems de las órdenes APROBADAS (aprobada, llegó, conforme), OC y OS, de las dos empresas.
# Los precios se comparan SIN IGV, solo dentro de la misma moneda y la misma U.M. (no se convierte de dólares a soles).
# "Ahorro potencial" = lo que se habría ahorrado si cada compra se hubiera hecho al mejor precio del periodo:
#     suma de (precio pagado - mejor precio) x cantidad.

ESTADOS_PRECIOS = ('APROBADA', 'RECIBIDA', 'CONFORME')
CLASES_PRECIOS = {'GENERAL': {'txt': 'Inventario General', 'corto': 'Producto', 'color': '#0B3D91', 'icono': 'bi-box-seam'},
                  'INSUMOS': {'txt': 'Insumos internos', 'corto': 'Insumo', 'color': '#0f766e', 'icono': 'bi-tools'},
                  'SERVICIO': {'txt': 'Servicios (OS)', 'corto': 'Servicio', 'color': '#7c3aed', 'icono': 'bi-gear'}}
VISTAS_PRECIOS = {'resumen': ('Resumen y gráficos', 'bi-graph-up'), 'productos': ('Por producto', 'bi-tags'),
                  'proveedores': ('Por proveedor', 'bi-building'), 'detalle': ('Detalle de compras', 'bi-list-ul')}
ORDENES_PRODUCTOS = {'ahorro': 'Mayor ahorro potencial', 'monto': 'Más comprado (monto)', 'compras': 'Más veces comprado',
                     'variacion': 'Último precio más caro que el mejor', 'proveedores': 'Más proveedores', 'nombre': 'Nombre (A-Z)'}
POR_PAGINA_PRECIOS = 50


def _oc_precios(emp_sel='TODAS', q='', desde=None, hasta=None, tipo='OC'):
    """(Compatibilidad) Precios por producto y proveedor de las órdenes aprobadas, agrupados por moneda y U.M."""
    f = _mp_filtros_por_defecto(empresa=emp_sel, q=q, desde=desde, hasta=hasta, tipo=tipo, clase='')
    lista = []
    for pr in _mp_productos(_mp_compras(f)):
        sg = {'moneda': pr['moneda'], 'simbolo': pr['simbolo'], 'um': pr['um'], 'filas': pr['ranking']}
        lista.append({'clave': pr['clave'], 'codigo': pr['codigo'], 'descripcion': pr['descripcion'], 'familia': pr['familia'],
                      'calidad': pr['calidad'], 'segmentos': [sg], 'n': pr['n'], 'n_prov': pr['n_prov']})
    juntos = {}
    for p in lista:          # un producto con compras en soles y en dólares: un solo bloque con dos segmentos
        if p['clave'] in juntos:
            juntos[p['clave']]['segmentos'] += p['segmentos']
            juntos[p['clave']]['n'] += p['n']
            juntos[p['clave']]['n_prov'] = len({f_['proveedor'] for sg in juntos[p['clave']]['segmentos'] for f_ in sg['filas']})
        else:
            juntos[p['clave']] = p
    salida = list(juntos.values())
    for p in salida:
        p['segmentos'].sort(key=lambda sg: (sg['moneda'] != 'PEN', sg['um']))
    salida.sort(key=lambda p: _clave_alfabetica(p['descripcion'] or ''))
    return salida


def _mp_filtros_por_defecto(**kw):
    f = {'empresa': 'TODAS', 'tipo': '', 'clase': '', 'moneda': '', 'prov': '', 'q': '', 'desde': None, 'hasta': None}
    f.update({k: v for k, v in kw.items() if k in f})
    return f


def _mp_filtros():
    """Filtros de la página (periodo, empresa, OC/OS, clase, moneda, proveedor, búsqueda)."""
    hoy = hora_peru().date()
    periodo = request.args.get('periodo', '12m')
    desde, hasta = _fecha_param('desde', None), _fecha_param('hasta', None)
    if not (desde or hasta):
        meses = {'3m': 3, '6m': 6, '12m': 12, '24m': 24}.get(periodo)
        desde = (hoy - timedelta(days=round(30.44 * meses))) if meses else None
        hasta = None
    else:
        periodo = 'rango'
    emp = (request.args.get('empresa') or 'TODAS').upper()
    tipo = (request.args.get('tipo') or '').upper()
    clase = (request.args.get('clase') or '').upper()
    moneda = (request.args.get('moneda') or '').upper()
    return {'empresa': emp if emp in EMPRESAS_OC else 'TODAS', 'tipo': tipo if tipo in TIPOS_OC else '',
            'clase': clase if clase in CLASES_PRECIOS else '', 'moneda': moneda if moneda in MONEDAS_OC else '',
            'prov': (request.args.get('prov') or '').strip()[:40], 'q': ' '.join((request.args.get('q') or '').split())[:80],
            'desde': desde, 'hasta': hasta, 'periodo': periodo}


def _mp_clase(o):
    return 'SERVICIO' if o.tipo == 'OS' else ('INSUMOS' if o.destino == 'INSUMOS' else 'GENERAL')


def _mp_condicion(o):
    if o.es_importacion:
        return 'IMPORTACIÓN'
    return ' '.join((o.condicion_pago or '').upper().split()) or 'SIN DATO'


def _mp_prov_clave(o):
    return f"p{o.oc_proveedor_id}" if o.oc_proveedor_id else 'r' + ((o.ruc or '') or (o.razon_social or '').upper())[:30]


def _mp_compras(f):
    """Una fila por ítem comprado (de las órdenes aprobadas que cumplen los filtros)."""
    ordenes = OrdenCompra.query.filter(OrdenCompra.estado.in_(list(ESTADOS_PRECIOS))).all()
    terminos = [t for t in _txt_sin_tildes(f.get('q') or '').upper().split() if t]
    compras = []
    for o in ordenes:
        if not o.fecha or (f['empresa'] != 'TODAS' and (o.empresa or 'ANCLAJES') != f['empresa']) or (f['tipo'] and o.tipo != f['tipo']):
            continue
        if (f['desde'] and o.fecha < f['desde']) or (f['hasta'] and o.fecha > f['hasta']) or (f['moneda'] and o.moneda != f['moneda']):
            continue
        clase = _mp_clase(o)
        if f['clase'] and clase != f['clase']:
            continue
        pk = _mp_prov_clave(o)
        if f['prov'] and pk != f['prov']:
            continue
        for i in o.items:
            if not i.precio_unitario or not i.cantidad:
                continue
            if terminos:
                texto = _txt_sin_tildes(f"{i.codigo or ''} {i.descripcion} {i.familia or ''} {i.calidad or ''} {i.tipo_servicio or ''} "
                                        f"{o.razon_social} {o.ruc or ''}").upper()
                if not all(t in texto for t in terminos):
                    continue
            base = (i.codigo or i.descripcion or '').upper().strip()
            clave = f"{clase}|{(i.tipo_servicio or '').upper()}|{base}" if clase == 'SERVICIO' else f"{clase}|{base}"
            precio = float(i.precio_unitario)
            compras.append({'orden': o, 'item': i, 'fecha': o.fecha, 'empresa': o.empresa or 'ANCLAJES', 'tipo': o.tipo, 'clase': clase,
                            'clave': clave, 'codigo': i.codigo or '', 'descripcion': i.descripcion or '', 'familia': i.familia or '',
                            'calidad': i.calidad or '', 'servicio': i.tipo_servicio or '', 'um': (i.um or '').upper(), 'moneda': o.moneda,
                            'precio': precio, 'con_igv': (i.precio_con_igv or round(precio * (1 + (0 if o.es_importacion else IGV_OC)), 4)),
                            'cantidad': float(i.cantidad), 'total': float(i.total or round(precio * i.cantidad, 2)),
                            'proveedor': o.razon_social or '', 'ruc': o.ruc or '', 'prov_clave': pk, 'importacion': bool(o.es_importacion),
                            'condicion': _mp_condicion(o), 'pagada': (o.estado_pago or '') == 'PAGADO'})
    return compras


def _mp_productos(compras):
    """Por producto (y por moneda + U.M.): mejor y último precio, promedio, ranking de proveedores, ahorro potencial e historial."""
    grupos = {}
    for c in compras:
        grupos.setdefault((c['clave'], c['moneda'], c['um']), []).append(c)
    productos = []
    for (clave, moneda, um), cs in grupos.items():
        mejor = min(cs, key=lambda c: (c['precio'], -c['fecha'].toordinal(), -c['orden'].id))
        ultimo = max(cs, key=lambda c: (c['fecha'], c['orden'].id, c['item'].id))
        cant = sum(c['cantidad'] for c in cs)
        monto = round(sum(c['total'] for c in cs), 2)
        ahorro = round(sum(max(c['precio'] - mejor['precio'], 0) * c['cantidad'] for c in cs), 2)
        provs = {}
        for c in cs:
            provs.setdefault(c['prov_clave'], []).append(c)
        ranking = []
        for pk, pc in provs.items():
            m = min(pc, key=lambda c: (c['precio'], -c['fecha'].toordinal()))
            u = max(pc, key=lambda c: (c['fecha'], c['orden'].id))
            ranking.append({'prov_clave': pk, 'proveedor': u['proveedor'], 'ruc': u['ruc'], 'mejor': m, 'ultimo': u, 'veces': len(pc),
                            'cantidad': round(sum(c['cantidad'] for c in pc), 4), 'monto': round(sum(c['total'] for c in pc), 2),
                            'promedio': round(sum(c['total'] for c in pc) / sum(c['cantidad'] for c in pc), 6) if sum(c['cantidad'] for c in pc) else 0})
        ranking.sort(key=lambda r: (r['mejor']['precio'], r['proveedor']))
        for r in ranking:
            r['dif_pct'] = round(100 * (r['mejor']['precio'] - mejor['precio']) / mejor['precio'], 1) if mejor['precio'] else 0
        ref = cs[-1]
        productos.append({
            'clave': clave, 'clase': ref['clase'], 'codigo': ultimo['codigo'], 'descripcion': ultimo['descripcion'], 'familia': ultimo['familia'],
            'calidad': ultimo['calidad'], 'servicio': ultimo['servicio'], 'um': um, 'moneda': moneda,
            'simbolo': MONEDAS_OC.get(moneda, {}).get('simbolo', moneda), 'mejor': mejor, 'ultimo': ultimo,
            'promedio': round(monto / cant, 6) if cant else 0, 'maximo': max(c['precio'] for c in cs), 'n': len(cs), 'n_prov': len(provs),
            'cantidad': round(cant, 4), 'monto': monto, 'ahorro': ahorro,
            'variacion': round(100 * (ultimo['precio'] - mejor['precio']) / mejor['precio'], 1) if mejor['precio'] else 0,
            'ranking': ranking,
            'historial': [{'f': c['fecha'].isoformat(), 'p': round(c['precio'], 6), 'prov': c['proveedor'], 'oc': _oc_codigo(c['orden'])}
                          for c in sorted(cs, key=lambda c: (c['fecha'], c['orden'].id))]})
    return productos


def _mp_proveedores(compras, productos):
    """Por proveedor: cuánto le compramos, cuántas veces tuvo el mejor precio y cuánto le pagamos de más."""
    mejores = {}
    for p in productos:
        if p['n_prov'] > 1:
            mejores[p['mejor']['prov_clave']] = mejores.get(p['mejor']['prov_clave'], 0) + 1
    mejor_de = {(p['clave'], p['moneda'], p['um']): p['mejor']['precio'] for p in productos}
    provs = {}
    for c in compras:
        pv = provs.setdefault(c['prov_clave'], {'prov_clave': c['prov_clave'], 'proveedor': c['proveedor'], 'ruc': c['ruc'], 'ordenes': set(),
                                                'items': 0, 'productos': set(), 'comparables': set(), 'PEN': 0.0, 'USD': 0.0,
                                                'sobre_PEN': 0.0, 'sobre_USD': 0.0, 'condiciones': {}, 'empresas': set(), 'clases': set(),
                                                'ultima': None, 'importacion': False, 'detalle': {}})
        o = c['orden']
        pv['ordenes'].add(o.id)
        pv['items'] += 1
        seg = (c['clave'], c['moneda'], c['um'])
        pv['productos'].add(seg)
        pv[c['moneda']] = pv.get(c['moneda'], 0) + c['total']
        pv['sobre_' + c['moneda']] = pv.get('sobre_' + c['moneda'], 0) + max(c['precio'] - mejor_de[seg], 0) * c['cantidad']
        pv['empresas'].add(c['empresa'])
        pv['clases'].add(c['clase'])
        pv['importacion'] = pv['importacion'] or c['importacion']
        if not pv['ultima'] or c['fecha'] > pv['ultima']:
            pv['ultima'], pv['proveedor'] = c['fecha'], c['proveedor']
        d = pv['detalle'].setdefault(seg, {'descripcion': c['descripcion'], 'codigo': c['codigo'], 'um': c['um'], 'moneda': c['moneda'],
                                           'simbolo': MONEDAS_OC.get(c['moneda'], {}).get('simbolo', ''), 'precios': [], 'mejor_global': mejor_de[seg]})
        d['precios'].append(c['precio'])
    # condición de pago: una vez por orden
    vistas = set()
    for c in compras:
        k = (c['prov_clave'], c['orden'].id)
        if k not in vistas:
            vistas.add(k)
            provs[c['prov_clave']]['condiciones'][c['condicion']] = provs[c['prov_clave']]['condiciones'].get(c['condicion'], 0) + 1
    n_comparables = {}
    for p in productos:
        if p['n_prov'] > 1:
            for r in p['ranking']:
                n_comparables[r['prov_clave']] = n_comparables.get(r['prov_clave'], 0) + 1
    salida = []
    for pk, pv in provs.items():
        cond = sorted(pv['condiciones'].items(), key=lambda x: -x[1])
        det = []
        for d in pv['detalle'].values():
            mi = min(d['precios'])
            det.append({**d, 'mejor_suyo': mi, 'ultimo_suyo': d['precios'][-1], 'veces': len(d['precios']),
                        'dif_pct': round(100 * (mi - d['mejor_global']) / d['mejor_global'], 1) if d['mejor_global'] else 0})
        det.sort(key=lambda d: (-d['dif_pct'], d['descripcion']))
        salida.append({**pv, 'n_ordenes': len(pv['ordenes']), 'n_productos': len(pv['productos']), 'n_mejor': mejores.get(pk, 0),
                       'n_comparables': n_comparables.get(pk, 0),
                       'pct_mejor': round(100 * mejores.get(pk, 0) / n_comparables[pk], 0) if n_comparables.get(pk) else None,
                       'PEN': round(pv['PEN'], 2), 'USD': round(pv['USD'], 2), 'sobre_PEN': round(pv['sobre_PEN'], 2), 'sobre_USD': round(pv['sobre_USD'], 2),
                       'condicion': cond[0][0] if cond else '', 'condiciones_txt': ', '.join(f"{k} ({v})" for k, v in cond),
                       'empresas_txt': ' y '.join(EMPRESAS_OC[e]['corto'] for e in sorted(pv['empresas'])),
                       'clases_txt': ', '.join(CLASES_PRECIOS[k]['corto'] for k in ('GENERAL', 'INSUMOS', 'SERVICIO') if k in pv['clases']),
                       'detalle': det})
    return salida


def _mp_resumen(compras, productos, proveedores):
    """Indicadores y series para los gráficos."""
    ordenes = {c['orden'].id: c['orden'] for c in compras}
    k = {'ordenes': len(ordenes), 'items': len(compras), 'productos': len(productos), 'proveedores': len(proveedores),
         'PEN': round(sum(c['total'] for c in compras if c['moneda'] == 'PEN'), 2), 'USD': round(sum(c['total'] for c in compras if c['moneda'] == 'USD'), 2),
         'ahorro_PEN': round(sum(p['ahorro'] for p in productos if p['moneda'] == 'PEN'), 2),
         'ahorro_USD': round(sum(p['ahorro'] for p in productos if p['moneda'] == 'USD'), 2),
         'multi': sum(1 for p in productos if p['n_prov'] > 1), 'subieron': sum(1 for p in productos if p['variacion'] > 0)}
    k['pct_multi'] = round(100 * k['multi'] / k['productos']) if k['productos'] else 0
    meses = {}
    for c in compras:
        m = c['fecha'].strftime('%Y-%m')
        meses.setdefault(m, {'PEN': 0.0, 'USD': 0.0})
        meses[m][c['moneda']] = meses[m].get(c['moneda'], 0) + c['total']
    nombres_mes = ['Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun', 'Jul', 'Ago', 'Set', 'Oct', 'Nov', 'Dic']
    por_mes = [{'mes': m, 'txt': f"{nombres_mes[int(m[5:]) - 1]} {m[2:4]}", 'PEN': round(v.get('PEN', 0), 2), 'USD': round(v.get('USD', 0), 2)}
               for m, v in sorted(meses.items())]

    def suma_por(campo, unico_por_orden=False):
        d, vistas = {}, set()
        for c in compras:
            if unico_por_orden:
                if c['orden'].id in vistas:
                    continue
                vistas.add(c['orden'].id)
                monto_pen = c['orden'].total if c['moneda'] == 'PEN' else 0
                monto_usd = c['orden'].total if c['moneda'] == 'USD' else 0
                x = d.setdefault(c[campo], {'clave': c[campo], 'PEN': 0.0, 'USD': 0.0, 'n': 0})
                x['PEN'] += monto_pen or 0
                x['USD'] += monto_usd or 0
                x['n'] += 1
            else:
                x = d.setdefault(c[campo], {'clave': c[campo], 'PEN': 0.0, 'USD': 0.0, 'n': 0})
                x[c['moneda']] += c['total']
                x['n'] += 1
        for x in d.values():
            x['PEN'], x['USD'] = round(x['PEN'], 2), round(x['USD'], 2)
        return sorted(d.values(), key=lambda x: (-x['PEN'], -x['USD']))
    condiciones = suma_por('condicion', unico_por_orden=True)
    empresas = suma_por('empresa')
    for e in empresas:
        e['txt'] = EMPRESAS_OC.get(e['clave'], {}).get('corto', e['clave'])
        e['color'] = EMPRESAS_OC.get(e['clave'], {}).get('color', '#64748b')
    clases = suma_por('clase')
    for e in clases:
        e['txt'] = CLASES_PRECIOS[e['clave']]['txt']
        e['color'] = CLASES_PRECIOS[e['clave']]['color']
    top_pen = sorted([p for p in proveedores if p['PEN']], key=lambda p: -p['PEN'])[:10]
    top_usd = sorted([p for p in proveedores if p['USD']], key=lambda p: -p['USD'])[:10]
    top_ahorro = sorted([p for p in productos if p['ahorro'] > 0], key=lambda p: -p['ahorro'])[:10]
    subidas = sorted([p for p in productos if p['variacion'] > 0 and p['n_prov'] >= 1 and p['n'] > 1], key=lambda p: -p['variacion'])[:8]
    pagadas = sum(1 for o in ordenes.values() if (o.estado_pago or '') == 'PAGADO')
    return {'k': k, 'por_mes': por_mes, 'condiciones': condiciones, 'empresas': empresas, 'clases': clases, 'top_pen': top_pen,
            'top_usd': top_usd, 'top_ahorro': top_ahorro, 'subidas': subidas, 'pagadas': pagadas, 'por_pagar': len(ordenes) - pagadas}


def _mp_orden_productos(productos, orden):
    claves = {'ahorro': lambda p: (-p['ahorro'], -p['monto']), 'monto': lambda p: (p['moneda'] != 'PEN', -p['monto']),
              'compras': lambda p: (-p['n'], -p['monto']), 'variacion': lambda p: (-p['variacion'], -p['ahorro']),
              'proveedores': lambda p: (-p['n_prov'], -p['ahorro']), 'nombre': lambda p: _clave_alfabetica(p['descripcion'] or '')}
    return sorted(productos, key=claves.get(orden, claves['ahorro']))


@app.route('/ordenes/precios')
@_ruta_oc()
def oc_precios():
    """Mejores precios: a qué precio y a quién compramos cada producto, insumo o servicio (con estadísticas, gráficos y Excel)."""
    f = _mp_filtros()
    vista = request.args.get('vista', 'resumen')
    vista = vista if vista in VISTAS_PRECIOS else 'resumen'
    orden = request.args.get('orden', 'ahorro')
    orden = orden if orden in ORDENES_PRODUCTOS else 'ahorro'
    compras = _mp_compras(f)
    productos = _mp_productos(compras)
    proveedores = _mp_proveedores(compras, productos)
    if request.args.get('exportar') == 'excel':
        return _mp_excel(f, compras, _mp_orden_productos(productos, orden), proveedores)
    res = _mp_resumen(compras, productos, proveedores)
    pagina = max(request.args.get('pagina', 1, type=int), 1)
    lista, total_lista = [], 0
    if vista == 'productos':
        todos = _mp_orden_productos(productos, orden)
        total_lista = len(todos)
        lista = todos[(pagina - 1) * POR_PAGINA_PRECIOS: pagina * POR_PAGINA_PRECIOS]
    elif vista == 'proveedores':
        todos = sorted(proveedores, key=lambda p: (-p['PEN'], -p['USD']))
        total_lista = len(todos)
        lista = todos[(pagina - 1) * POR_PAGINA_PRECIOS: pagina * POR_PAGINA_PRECIOS]
    elif vista == 'detalle':
        mejor_de = {(p['clave'], p['moneda'], p['um']): p['mejor']['precio'] for p in productos}
        todos = sorted(compras, key=lambda c: (c['fecha'], c['orden'].id), reverse=True)
        for c in todos:
            mb = mejor_de[(c['clave'], c['moneda'], c['um'])]
            c['dif_pct'] = round(100 * (c['precio'] - mb) / mb, 1) if mb else 0
        total_lista = len(todos)
        lista = todos[(pagina - 1) * 100: pagina * 100]
    por_pag = 100 if vista == 'detalle' else POR_PAGINA_PRECIOS
    opciones_prov = sorted({(c['prov_clave'], c['proveedor']) for c in _mp_compras({**f, 'prov': '', 'q': ''})}, key=lambda x: _clave_alfabetica(x[1]))
    return render_template('oc_precios.html', f=f, vista=vista, vistas=VISTAS_PRECIOS, orden=orden, ordenes_productos=ORDENES_PRODUCTOS,
                           res=res, lista=lista, total_lista=total_lista, pagina=pagina, paginas=max(1, -(-total_lista // por_pag)),
                           clases=CLASES_PRECIOS, tipos=TIPOS_OC, empresas_oc=EMPRESAS_OC, monedas=MONEDAS_OC, igv=IGV_OC,
                           opciones_prov=opciones_prov, seccion='precios', puede_ver_oc=session.get('role') in ROLES_OC_VEN,
                           hoy=hora_peru().date())


def _mp_excel(f, compras, productos, proveedores):
    """Excel profesional: resumen con gráficos, mejores precios, ranking de proveedores por producto, por proveedor y detalle."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.chart.label import DataLabelList
    wb = Workbook()
    res = _mp_resumen(compras, productos, proveedores)
    m2, m4 = '#,##0.00', '#,##0.00##'
    filtro_txt = (f"{'Las dos empresas' if f['empresa'] == 'TODAS' else EMPRESAS_OC[f['empresa']]['nombre']} · "
                  f"{'OC y OS' if not f['tipo'] else TIPOS_OC[f['tipo']]} · {CLASES_PRECIOS[f['clase']]['txt'] if f['clase'] else 'Productos, insumos y servicios'} · "
                  f"{('del ' + f['desde'].strftime('%d/%m/%Y')) if f['desde'] else 'desde el inicio'} "
                  f"{('al ' + f['hasta'].strftime('%d/%m/%Y')) if f['hasta'] else 'a hoy'}"
                  f"{' · solo ' + MONEDAS_OC[f['moneda']]['nombre'] if f['moneda'] else ''}{' · búsqueda: ' + f['q'] if f['q'] else ''}")
    nota = f"Órdenes aprobadas · precios SIN IGV · soles y dólares por separado · {filtro_txt}"
    # ---- Resumen ----
    ws = wb.active
    ws.title = 'Resumen'
    azul, gris = PatternFill('solid', fgColor='0B3D91'), PatternFill('solid', fgColor='F1F5F9')
    ws['A1'] = 'MEJORES PRECIOS DE COMPRA'
    ws['A1'].font = Font(bold=True, size=16, color='0B3D91')
    ws['A2'] = nota
    ws['A2'].font = Font(italic=True, color='475569')
    ws['A3'] = f"Generado el {hora_peru().strftime('%d/%m/%Y %H:%M')} por {session.get('nombre', '')}"
    ws['A3'].font = Font(color='64748B', size=9)
    k = res['k']
    indicadores = [('Órdenes aprobadas', k['ordenes'], None), ('Ítems comprados', k['items'], None), ('Productos / insumos / servicios', k['productos'], None),
                   ('Proveedores', k['proveedores'], None), ('Comprado en soles (sin IGV)', k['PEN'], m2), ('Comprado en dólares (sin IGV)', k['USD'], m2),
                   ('Ahorro potencial en soles', k['ahorro_PEN'], m2), ('Ahorro potencial en dólares', k['ahorro_USD'], m2),
                   ('Productos con 2 o más proveedores', k['multi'], None), ('Productos cuyo último precio es mayor al mejor', k['subieron'], None)]
    ws['A5'], ws['B5'] = 'Indicador', 'Valor'
    for c in (ws['A5'], ws['B5']):
        c.font, c.fill = Font(bold=True, color='FFFFFF'), azul
    for n, (t, v, fm) in enumerate(indicadores, start=6):
        ws.cell(row=n, column=1, value=t).fill = gris
        c = ws.cell(row=n, column=2, value=v)
        c.font = Font(bold=True)
        if fm:
            c.number_format = fm
    ws.cell(row=17, column=1, value='"Ahorro potencial": lo que se habría ahorrado comprando todo al mejor precio del periodo (mismo producto, moneda y U.M.).').font = Font(italic=True, size=9, color='64748B')
    ws.column_dimensions['A'].width, ws.column_dimensions['B'].width = 48, 18

    def tabla(fila, col, titulo, cab, datos, fmts=None):
        ws.cell(row=fila, column=col, value=titulo).font = Font(bold=True, color='0B3D91', size=12)
        for j, t in enumerate(cab):
            c = ws.cell(row=fila + 1, column=col + j, value=t)
            c.font, c.fill = Font(bold=True, color='FFFFFF'), azul
        for i, d in enumerate(datos, start=fila + 2):
            for j, v in enumerate(d):
                c = ws.cell(row=i, column=col + j, value=v)
                if fmts and j in fmts:
                    c.number_format = fmts[j]
        return fila + 1, fila + 1 + len(datos)
    for col, w in (('D', 26), ('E', 16), ('F', 16), ('G', 10)):
        ws.column_dimensions[col].width = w
    # gasto por mes
    c0, c1 = tabla(5, 4, 'Compras por mes (sin IGV)', ['Mes', 'Soles (S/)', 'Dólares (US$)'], [[x['txt'], x['PEN'], x['USD']] for x in res['por_mes']], {1: m2, 2: m2})
    if res['por_mes']:
        g = BarChart()
        g.type, g.grouping, g.title, g.y_axis.title = 'col', 'clustered', 'Compras por mes (sin IGV)', 'Monto'
        g.add_data(Reference(ws, min_col=5, max_col=6, min_row=c0, max_row=c1), titles_from_data=True)
        g.set_categories(Reference(ws, min_col=4, min_row=c0 + 1, max_row=c1))
        g.height, g.width = 7.5, 17
        ws.add_chart(g, 'I5')
    fila = max(c1 + 3, 20)
    # top proveedores
    tp = res['top_pen'] or res['top_usd']
    mon = 'PEN' if res['top_pen'] else 'USD'
    p0, p1 = tabla(fila, 4, f"Top proveedores ({MONEDAS_OC[mon]['simbolo']})", ['Proveedor', f"Monto {MONEDAS_OC[mon]['simbolo']}", 'Órdenes'],
                   [[p['proveedor'][:40], p[mon], p['n_ordenes']] for p in tp], {1: m2})
    if tp:
        g = BarChart()
        g.type, g.title = 'bar', f"Top proveedores por monto ({MONEDAS_OC[mon]['simbolo']})"
        g.add_data(Reference(ws, min_col=5, min_row=p0, max_row=p1), titles_from_data=True)
        g.set_categories(Reference(ws, min_col=4, min_row=p0 + 1, max_row=p1))
        g.legend = None
        g.height, g.width = 8, 17
        ws.add_chart(g, f'I{fila}')
    fila = max(p1 + 3, fila + 18)
    # condición de pago
    q0, q1 = tabla(fila, 4, 'Condición de pago (órdenes)', ['Condición', 'Órdenes', 'Soles (S/)', 'Dólares (US$)'],
                   [[x['clave'], x['n'], x['PEN'], x['USD']] for x in res['condiciones']], {2: m2, 3: m2})
    if res['condiciones']:
        g = PieChart()
        g.title = 'Condición de pago (n° de órdenes)'
        g.add_data(Reference(ws, min_col=5, min_row=q0, max_row=q1), titles_from_data=True)
        g.set_categories(Reference(ws, min_col=4, min_row=q0 + 1, max_row=q1))
        g.dataLabels = DataLabelList()
        g.dataLabels.showPercent = True
        g.height, g.width = 7.5, 12
        ws.add_chart(g, f'I{fila}')
    fila = max(q1 + 3, fila + 17)
    tabla(fila, 4, 'Por empresa', ['Empresa', 'Ítems', 'Soles (S/)', 'Dólares (US$)'], [[x['txt'], x['n'], x['PEN'], x['USD']] for x in res['empresas']], {2: m2, 3: m2})
    fila += len(res['empresas']) + 4
    tabla(fila, 4, 'Por tipo de compra', ['Tipo', 'Ítems', 'Soles (S/)', 'Dólares (US$)'], [[x['txt'], x['n'], x['PEN'], x['USD']] for x in res['clases']], {2: m2, 3: m2})
    # ---- Mejores precios ----
    _xlsx_hoja(wb, 'Mejores precios', ['Tipo', 'Código', 'Producto / servicio', 'Familia / categoría', 'Calidad', 'U.M.', 'Moneda',
                                       'Mejor precio sin IGV', 'Mejor precio con IGV', 'Proveedor del mejor precio', 'Fecha del mejor', 'Orden del mejor',
                                       'Último precio sin IGV', 'Proveedor del último', 'Fecha del último', 'Precio promedio', 'Precio más alto',
                                       'Último vs mejor (%)', 'Veces comprado', 'N° proveedores', 'Cantidad comprada', 'Monto comprado (sin IGV)', 'Ahorro potencial'],
               [[CLASES_PRECIOS[p['clase']]['corto'] + (f" · {p['servicio']}" if p['servicio'] else ''), p['codigo'], p['descripcion'], p['familia'], p['calidad'],
                 p['um'], p['moneda'], p['mejor']['precio'], None if p['mejor']['importacion'] else p['mejor']['con_igv'], p['mejor']['proveedor'],
                 p['mejor']['fecha'], _oc_codigo(p['mejor']['orden']), p['ultimo']['precio'], p['ultimo']['proveedor'], p['ultimo']['fecha'],
                 p['promedio'], p['maximo'], p['variacion'] / 100, p['n'], p['n_prov'], p['cantidad'], p['monto'], p['ahorro']] for p in productos],
               formatos={7: m4, 8: m4, 10: 'dd/mm/yyyy', 12: m4, 14: 'dd/mm/yyyy', 15: m4, 16: m4, 17: '0.0%', 20: '#,##0.##', 21: m2, 22: m2}, nota=nota)
    # ---- Ranking por producto ----
    filas = []
    for p in productos:
        for n, r in enumerate(p['ranking'], start=1):
            filas.append([p['codigo'], p['descripcion'], p['um'], p['moneda'], n, r['proveedor'], r['ruc'], r['mejor']['precio'],
                          None if r['mejor']['importacion'] else r['mejor']['con_igv'], r['mejor']['fecha'], _oc_codigo(r['mejor']['orden']),
                          r['ultimo']['precio'], r['ultimo']['fecha'], r['promedio'], r['veces'], r['cantidad'], r['dif_pct'] / 100])
    _xlsx_hoja(wb, 'Ranking por producto', ['Código', 'Producto / servicio', 'U.M.', 'Moneda', 'Puesto', 'Proveedor', 'RUC / Tax ID', 'Mejor precio sin IGV',
                                            'Mejor precio con IGV', 'Fecha del mejor', 'Orden del mejor', 'Último precio sin IGV', 'Fecha del último',
                                            'Precio promedio', 'Veces', 'Cantidad', '% más caro que el mejor'],
               filas, formatos={7: m4, 8: m4, 9: 'dd/mm/yyyy', 11: m4, 12: 'dd/mm/yyyy', 13: m4, 15: '#,##0.##', 16: '0.0%'}, nota=nota)
    # ---- Por proveedor ----
    _xlsx_hoja(wb, 'Por proveedor', ['Proveedor', 'RUC / Tax ID', 'Del exterior', 'Órdenes', 'Ítems', 'Productos', 'Comprado S/ (sin IGV)', 'Comprado US$ (sin IGV)',
                                     'Veces con el mejor precio', 'Productos comparables', '% mejor precio', 'Pagado de más S/', 'Pagado de más US$',
                                     'Condición de pago más usada', 'Condiciones de pago', 'Empresas', 'Tipo', 'Última compra'],
               [[p['proveedor'], p['ruc'], 'Sí' if p['importacion'] else '', p['n_ordenes'], p['items'], p['n_productos'], p['PEN'], p['USD'], p['n_mejor'],
                 p['n_comparables'], (p['pct_mejor'] / 100) if p['pct_mejor'] is not None else None, p['sobre_PEN'], p['sobre_USD'], p['condicion'],
                 p['condiciones_txt'], p['empresas_txt'], p['clases_txt'], p['ultima']] for p in sorted(proveedores, key=lambda p: (-p['PEN'], -p['USD']))],
               formatos={6: m2, 7: m2, 10: '0%', 11: m2, 12: m2, 17: 'dd/mm/yyyy'}, nota=nota)
    # ---- Detalle ----
    mejor_de = {(p['clave'], p['moneda'], p['um']): p['mejor']['precio'] for p in productos}
    _xlsx_hoja(wb, 'Detalle de compras', ['Fecha', 'Empresa', 'Orden', 'Tipo', 'Estado', 'Proveedor', 'RUC / Tax ID', 'Código', 'Producto / servicio', 'Familia',
                                          'Calidad', 'Cantidad', 'U.M.', 'Moneda', 'P. unit. sin IGV', 'P. unit. con IGV', 'Total sin IGV', 'Mejor precio del periodo',
                                          'vs. el mejor (%)', 'Condición de pago', 'Pagada', 'Importación'],
               [[c['fecha'], EMPRESAS_OC[c['empresa']]['corto'], _oc_codigo(c['orden']), CLASES_PRECIOS[c['clase']]['corto'], ESTADOS_OC[c['orden'].estado]['txt'],
                 c['proveedor'], c['ruc'], c['codigo'], c['descripcion'], c['familia'], c['calidad'], c['cantidad'], c['um'], c['moneda'], c['precio'],
                 None if c['importacion'] else c['con_igv'], c['total'], mejor_de[(c['clave'], c['moneda'], c['um'])],
                 ((c['precio'] - mejor_de[(c['clave'], c['moneda'], c['um'])]) / mejor_de[(c['clave'], c['moneda'], c['um'])]) if mejor_de[(c['clave'], c['moneda'], c['um'])] else 0,
                 c['condicion'], 'Sí' if c['pagada'] else 'No', 'Sí' if c['importacion'] else ''] for c in sorted(compras, key=lambda c: (c['fecha'], c['orden'].id), reverse=True)],
               formatos={0: 'dd/mm/yyyy', 11: '#,##0.##', 14: m4, 15: m4, 16: m2, 17: m4, 18: '0.0%'}, nota=nota)
    registrar_log(f"Exportó Mejores precios a Excel ({len(productos)} producto(s), {len(compras)} compra(s))", 'bi-file-earmark-excel', 'text-success')
    return _xlsx_respuesta(wb, f"mejores_precios_{hora_peru().strftime('%Y%m%d_%H%M')}.xlsx")


@app.route('/reporte_precios_proveedores')
def reporte_precios_proveedores():
    """El menú antiguo "Mejores Precios" ahora abre el análisis de compras (las órdenes aprobadas)."""
    if session.get('user_id') is None:
        return redirect(url_for('login'))
    if session.get('role') not in ROLES_OC_VEN:
        return "Acceso denegado", 403
    return redirect(url_for('oc_precios', **request.args))
