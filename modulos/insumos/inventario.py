# flake8: noqa
# pyright: reportUndefinedVariable=false
# =====================================================================================
# INSUMOS INTERNOS · inventario
# inventario por empresa, importar desde Excel, maestro de insumos y proveedores de insumos
# Este archivo NO se importa: app.py lo carga en su lugar (lista MODULOS_EN_ARCHIVOS), como si
# su código estuviera escrito ahí. Usa directamente los modelos, variables y funciones de app.py
# y de los archivos que se cargan antes. Ver MAPA.md.
# =====================================================================================

# ---------------------------- INVENTARIO ----------------------------
@app.route('/insumos')
@_ruta_insumos()
def insumos_inventario():
    empresa = _empresa_param()
    session['insumos_empresa'] = empresa
    q = ' '.join((request.args.get('q') or '').split())
    categoria = request.args.get('categoria', '')
    bajo = request.args.get('bajo') == '1'
    ver = request.args.get('ver', 'activos')
    insumos = Insumo.query.order_by(Insumo.codigo).all()
    stocks = {s.insumo_id: s for s in InsumoStock.query.filter_by(empresa=empresa).all()}
    filas = []
    for i in insumos:
        if ver == 'activos' and not i.activo:
            continue
        if categoria and (i.categoria or '') != categoria:
            continue
        if q and not all(p in f'{i.codigo} {i.descripcion}'.upper() for p in q.upper().split()):
            continue
        s = stocks.get(i.id)
        stock = float(s.stock_actual or 0) if s else 0
        minimo = float(s.stock_minimo or 0) if s else 0
        if bajo and not (minimo > 0 and stock <= minimo):
            continue
        q_pen, c_pen = _bolsa(s, 'PEN') if s else (0.0, 0.0)
        q_usd, c_usd = _bolsa(s, 'USD') if s else (0.0, 0.0)
        filas.append({'i': i, 's': s, 'stock': stock, 'minimo': minimo, 'costo': c_pen, 'costo_usd': c_usd, 'stock_pen': q_pen, 'stock_usd': q_usd,
                      'valor': q_pen * c_pen, 'valor_usd': q_usd * c_usd, 'bajo': minimo > 0 and stock <= minimo})
    ahora = hora_peru()
    inicio_mes = datetime(ahora.year, ahora.month, 1)
    sal_mes = InsumoSalida.query.filter(InsumoSalida.empresa == empresa, InsumoSalida.estado == 'REGISTRADA', InsumoSalida.fecha >= inicio_mes).all()
    gasto_mes = sum(x.total_costo or 0 for x in sal_mes)
    gasto_mes_usd = sum(x.total_costo_usd or 0 for x in sal_mes)
    todos = [{'stock': float(s.stock_actual or 0), 'min': float(s.stock_minimo or 0),
              'valor': _bolsa(s, 'PEN')[0] * _bolsa(s, 'PEN')[1], 'valor_usd': _bolsa(s, 'USD')[0] * _bolsa(s, 'USD')[1]}
             for s in stocks.values()]
    kpi = {'items': sum(1 for t in todos if t['stock'] > 0), 'valor': sum(t['valor'] for t in todos), 'valor_usd': sum(t['valor_usd'] for t in todos),
           'bajo': sum(1 for t in todos if t['min'] > 0 and t['stock'] <= t['min']), 'gasto_mes': gasto_mes, 'gasto_mes_usd': gasto_mes_usd}
    return render_template('insumos_inventario.html', empresa=empresa, empresas=EMPRESAS_INSUMOS, filas=filas, kpi=kpi,
                           q=q, categoria=categoria, bajo=bajo, ver=ver, categorias=_categorias_insumos(),
                           almacenes=_catalogo_insumos('INSUMO_ALMACEN'), puede_editar=session.get('role') in ROLES_INSUMOS_EDITAN,
                           puede_ajustar=session.get('role') == 'admin', oc_pendientes=_ordenes_pendientes_insumos(empresa) if _oc_ok else [],
                           seccion='inventario')


@app.route('/insumos/stock/<int:insumo_id>/<empresa>', methods=['POST'])
@_ruta_insumos(editar=True, como_json=True)
def insumos_stock_config(insumo_id, empresa):
    """Stock mínimo y almacén/ubicación de un insumo en una empresa."""
    empresa = empresa.upper()
    if empresa not in EMPRESAS_INSUMOS or not Insumo.query.get(insumo_id):
        return {'status': 'error', 'msg': 'Datos inválidos'}
    minimo = _num(request.form.get('stock_minimo'))
    if minimo is None or minimo < 0:
        return {'status': 'error', 'msg': 'El stock mínimo debe ser un número mayor o igual a 0.'}
    st = _stock_de(insumo_id, empresa)
    st.stock_minimo = minimo
    st.almacen = ' '.join((request.form.get('almacen') or '').split())[:100] or None
    db.session.commit()
    return {'status': 'success'}


# ---------------------------- IMPORTAR INSUMOS DESDE EXCEL (por inventario) ----------------------------
COLUMNAS_IMPORT_INSUMOS = {      # cabecera normalizada -> campo
    'CODIGO': 'codigo', 'COD': 'codigo', 'SKU': 'codigo',
    'DESCRIPCION': 'descripcion', 'INSUMO': 'descripcion', 'PRODUCTO': 'descripcion',
    'CATEGORIA': 'categoria',
    'STOCK': 'stock', 'STOCK ACTUAL': 'stock', 'CANTIDAD': 'stock',
    'MINIMO': 'minimo', 'STOCK MINIMO': 'minimo', 'MIN': 'minimo',
    'COSTO PROM': 'costo', 'COSTO PROMEDIO': 'costo', 'COSTO': 'costo', 'COSTO UNITARIO': 'costo', 'COSTO PROM S': 'costo',
    'UNIDAD': 'unidad', 'UM': 'unidad', 'U M': 'unidad', 'UNIDAD DE MEDIDA': 'unidad',
}
MAX_FILAS_IMPORT_INSUMOS = 3000


def _norm_cab(v):
    """'Costo prom.' -> 'COSTO PROM' ; 'U.M.' -> 'U M' ; 'Código' -> 'CODIGO'."""
    return ' '.join(re.sub(r'[^A-Z0-9 ]', ' ', _txt_sin_tildes(str(v or '')).upper()).split())


def _leer_excel_insumos(archivo):
    """Lee la primera hoja: busca la fila de cabeceras (Código, Descripción, Categoría, Stock; opcionales Mínimo, Costo prom., Unidad).
    Devuelve (filas [{fila, codigo, descripcion, categoria, stock, minimo, costo, unidad}], error)."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(archivo, read_only=True, data_only=True)
    except Exception:
        return [], 'No se pudo leer el archivo: súbelo en formato Excel (.xlsx).'
    ws = wb.worksheets[0]
    filas_xl = list(ws.iter_rows(values_only=True))
    mapa, fila_cab = None, None
    for n, fila in enumerate(filas_xl[:15]):
        m = {}
        for j, v in enumerate(fila or []):
            campo = COLUMNAS_IMPORT_INSUMOS.get(_norm_cab(v))
            if campo and campo not in m:
                m[campo] = j
        if 'descripcion' in m and 'stock' in m:
            mapa, fila_cab = m, n
            break
    if not mapa:
        return [], 'No encontré las columnas. La primera fila debe tener: Código, Descripción, Categoría, Stock (y si quieres Mínimo, Costo prom., Unidad).'
    filas = []
    for n, fila in enumerate(filas_xl[fila_cab + 1:], start=fila_cab + 2):
        fila = list(fila or [])
        val = lambda campo: fila[mapa[campo]] if campo in mapa and mapa[campo] < len(fila) else None
        txt = lambda campo, largo: ' '.join(str(val(campo) if val(campo) is not None else '').split())[:largo]
        if not any(x not in (None, '') for x in fila):
            continue
        filas.append({'fila': n, 'codigo': txt('codigo', 50).upper(), 'descripcion': txt('descripcion', 300).upper(),
                      'categoria': txt('categoria', 100).upper(), 'unidad': txt('unidad', 30).upper(),
                      'stock': val('stock'), 'minimo': val('minimo'), 'costo': val('costo')})
        if len(filas) > MAX_FILAS_IMPORT_INSUMOS:
            return [], f'El archivo tiene más de {MAX_FILAS_IMPORT_INSUMOS} filas: divídelo en partes.'
    if not filas:
        return [], 'El archivo no tiene filas con datos debajo de las cabeceras.'
    return filas, None


def _analizar_import_insumos(filas, empresa):
    """Compara cada fila con el maestro y el stock de la empresa: qué se crea, qué cambia y qué errores hay. No guarda nada."""
    insumos = Insumo.query.all()
    por_cod = {(i.codigo or '').upper(): i for i in insumos}
    norm = lambda t: ' '.join(_clave_alfabetica(t or '').upper().split())
    por_desc = {}
    for i in insumos:
        por_desc.setdefault(norm(i.descripcion), i)
    stocks = {st.insumo_id: st for st in InsumoStock.query.filter_by(empresa=empresa).all()}
    unidades = {(v.valor or '').upper() for v in _catalogo_insumos('INSUMO_UNIDAD', incluir_inactivos=True)}
    vistos_cod, vistos_desc = {}, {}
    res = []
    for f in filas:
        r = {**f, 'errores': [], 'avisos': [], 'insumo_id': None, 'nuevo': False}
        r['stock'] = _num(f.get('stock')) if f.get('stock') not in (None, '') else None
        r['minimo'] = _num(f.get('minimo')) if f.get('minimo') not in (None, '') else None
        r['costo'] = _num(f.get('costo')) if f.get('costo') not in (None, '') else None
        if len(r['descripcion']) < 3:
            r['errores'].append('Falta la descripción.')
        if r['stock'] is None or r['stock'] < 0:
            r['errores'].append('El stock debe ser un número mayor o igual a 0.')
        if f.get('minimo') not in (None, '') and (r['minimo'] is None or r['minimo'] < 0):
            r['errores'].append('El mínimo no es válido.')
        if f.get('costo') not in (None, '') and (r['costo'] is None or r['costo'] < 0):
            r['errores'].append('El costo no es válido.')
        ins = por_cod.get(r['codigo']) if r['codigo'] else None
        if not ins and r['descripcion']:
            ins = por_desc.get(norm(r['descripcion']))
            if ins and r['codigo'] and ins.codigo.upper() != r['codigo']:
                r['avisos'].append(f'Ya existe con el código {ins.codigo}: se usa ese.')
        if ins and r['codigo'] and r['codigo'] == ins.codigo.upper() and r['descripcion'] and norm(ins.descripcion) != norm(r['descripcion']):
            r['avisos'].append(f'En el maestro se llama "{ins.descripcion}": se mantiene ese nombre.')
        clave_c, clave_d = (ins.codigo.upper() if ins else r['codigo']), norm(r['descripcion'])
        if (clave_c and clave_c in vistos_cod) or (not ins and clave_d in vistos_desc):
            r['errores'].append(f"Repetido: ya está en la fila {vistos_cod.get(clave_c) or vistos_desc.get(clave_d)}.")
        if clave_c:
            vistos_cod.setdefault(clave_c, f['fila'])
        vistos_desc.setdefault(clave_d, f['fila'])
        if ins:
            r['insumo_id'], r['codigo_final'], r['descripcion_final'] = ins.id, ins.codigo, ins.descripcion
            r['categoria_final'] = ins.categoria or r['categoria'] or ''
            if not ins.activo:
                r['avisos'].append('Estaba desactivado: se vuelve a activar.')
        else:
            r['nuevo'] = True
            if r['codigo'] and r['codigo'] in por_cod:
                r['errores'].append('Ese código ya existe.')
            r['codigo_final'], r['descripcion_final'], r['categoria_final'] = r['codigo'] or '(se genera)', r['descripcion'], r['categoria']
            if not r['unidad']:
                r['unidad'] = 'UND'
                r['avisos'].append('Sin unidad: se pone UND.')
        st = stocks.get(ins.id) if ins else None
        r['stock_antes'] = round(float(st.stock_actual or 0), 6) if st else 0.0
        r['costo_antes'] = round(float(st.costo_promedio or 0), 6) if st else 0.0
        r['minimo_antes'] = round(float(st.stock_minimo or 0), 6) if st else 0.0
        r['dif'] = round((r['stock'] or 0) - r['stock_antes'], 6) if r['stock'] is not None else 0
        if r['stock'] and not r['costo'] and not r['costo_antes']:
            r['avisos'].append('Sin costo: el stock entra con costo 0 (puedes corregirlo después).')
        if r['unidad'] and r['unidad'] not in unidades and r['nuevo']:
            r['avisos'].append(f"Unidad {r['unidad']} nueva: se agrega a la lista.")
        res.append(r)
    resumen = {'filas': len(res), 'errores': sum(1 for r in res if r['errores']), 'nuevos': sum(1 for r in res if r['nuevo'] and not r['errores']),
               'cambian_stock': sum(1 for r in res if not r['errores'] and abs(r['dif']) > 1e-9),
               'sube': sum(1 for r in res if not r['errores'] and r['dif'] > 1e-9), 'baja': sum(1 for r in res if not r['errores'] and r['dif'] < -1e-9)}
    return res, resumen


@app.route('/insumos/importar', methods=['GET', 'POST'])
@_ruta_insumos(editar=True)
def insumos_importar():
    """Carga desde Excel al inventario de UNA empresa: crea en el maestro lo que no existe y deja el stock como dice el archivo
    (la diferencia queda en el Kardex como 'Carga desde Excel', con quién y cuándo). Primero muestra una vista previa."""
    empresa = _empresa_param()
    if request.method == 'GET':
        return render_template('insumos_importar.html', empresa=empresa, empresas=EMPRESAS_INSUMOS, filas=None, resumen=None,
                               seccion='inventario')
    request.max_content_length = LIMITE_PETICION_MOVIMIENTO
    paso = request.form.get('paso', 'previa')
    if paso == 'previa':
        archivo = request.files.get('archivo')
        if not archivo or not archivo.filename:
            flash('⛔ Elige el archivo Excel (.xlsx).')
            return redirect(url_for('insumos_importar', empresa=empresa))
        if not archivo.filename.lower().endswith(('.xlsx', '.xlsm')):
            flash('⛔ El archivo debe ser Excel (.xlsx). Si lo tienes en .xls o .csv, ábrelo y "Guardar como" Libro de Excel (.xlsx).')
            return redirect(url_for('insumos_importar', empresa=empresa))
        filas, error = _leer_excel_insumos(io.BytesIO(archivo.read()))
        if error:
            flash('⛔ ' + error)
            return redirect(url_for('insumos_importar', empresa=empresa))
        res, resumen = _analizar_import_insumos(filas, empresa)
        datos = json.dumps([{k: f[k] for k in ('fila', 'codigo', 'descripcion', 'categoria', 'unidad', 'stock', 'minimo', 'costo')} for f in filas],
                           ensure_ascii=False, default=str)
        return render_template('insumos_importar.html', empresa=empresa, empresas=EMPRESAS_INSUMOS, filas=res, resumen=resumen, datos=datos,
                               nombre_archivo=secure_filename(archivo.filename), seccion='inventario')
    # paso 'confirmar': se vuelve a validar contra lo que hay AHORA (por si alguien movió el stock mientras tanto)
    try:
        filas = json.loads(request.form.get('datos') or '[]')
        assert isinstance(filas, list) and all(isinstance(f, dict) for f in filas)
    except (ValueError, AssertionError):
        flash('⛔ Los datos de la vista previa no son válidos: vuelve a subir el archivo.')
        return redirect(url_for('insumos_importar', empresa=empresa))
    res, resumen = _analizar_import_insumos(filas, empresa)
    if resumen['errores'] or not res:
        flash(f"⛔ No se importó nada: hay {resumen['errores']} fila(s) con errores. Corrígelas en el Excel y vuelve a subirlo.")
        return redirect(url_for('insumos_importar', empresa=empresa))
    ahora = hora_peru()
    nombre_arch = ' '.join((request.form.get('nombre_archivo') or '').split())[:60]
    ref = f"Carga desde Excel{(' · ' + nombre_arch) if nombre_arch else ''}"[:120]
    cats = {(v.valor or '').upper() for v in _catalogo_insumos('INSUMO_CATEGORIA', incluir_inactivos=True)}
    unis = {(v.valor or '').upper() for v in _catalogo_insumos('INSUMO_UNIDAD', incluir_inactivos=True)}
    creados = movs = 0
    try:
        for r in res:
            if r['categoria'] and r['categoria'] not in cats:
                db.session.add(CatalogoValor(tipo='INSUMO_CATEGORIA', valor=r['categoria'], activo=True, es_predeterminado=False, creado_por_id=session.get('user_id')))
                cats.add(r['categoria'])
            if r['nuevo'] and r['unidad'] and r['unidad'] not in unis:
                db.session.add(CatalogoValor(tipo='INSUMO_UNIDAD', valor=r['unidad'], activo=True, es_predeterminado=False, creado_por_id=session.get('user_id')))
                unis.add(r['unidad'])
            if r['nuevo']:
                ins = Insumo(codigo=r['codigo'] or f'TMP-{uuid.uuid4().hex[:8]}', descripcion=r['descripcion'], categoria=r['categoria'] or None,
                             unidad=r['unidad'] or 'UND', activo=True, creado_por=session.get('nombre', ''), fecha_creacion=ahora)
                db.session.add(ins)
                db.session.flush()
                if not r['codigo']:
                    ins.codigo = f'INS-{ins.id:05d}'
                creados += 1
            else:
                ins = Insumo.query.get(r['insumo_id'])
                if not ins.activo:
                    ins.activo = True
                if r['categoria'] and not ins.categoria:
                    ins.categoria = r['categoria']
            st = _stock_de(ins.id, empresa)
            if r['minimo'] is not None:
                st.stock_minimo = r['minimo']
            dif = round(r['stock'] - float(st.stock_actual or 0), 6)
            if dif > 1e-9:
                costo = r['costo'] if r['costo'] is not None else float(st.costo_promedio or 0)
                antes, _ = _entrada_a_stock(st, dif, costo, 'PEN')
                db.session.add(InsumoMovimiento(empresa=empresa, insumo_id=ins.id, fecha=ahora, tipo='ENTRADA', concepto='IMPORTACION', cantidad=dif,
                                                costo_unitario=round(costo, 6), costo_total=round(dif * costo, 2), stock_antes=antes,
                                                stock_despues=st.stock_actual, **_datos_mov_moneda(st, 'PEN'), referencia=ref,
                                                observacion='Stock según el Excel (inventario inicial o conteo)', user_id=session.get('user_id'),
                                                fecha_registro=ahora))
                st.ultimo_ingreso = ahora
                movs += 1
            elif dif < -1e-9:
                antes, partes = _salida_por_monedas(st, -dif)
                for mon, cant_p, costo in partes:
                    db.session.add(InsumoMovimiento(empresa=empresa, insumo_id=ins.id, fecha=ahora, tipo='SALIDA', concepto='IMPORTACION', cantidad=cant_p,
                                                    costo_unitario=round(costo, 6), costo_total=round(cant_p * costo, 2), stock_antes=antes,
                                                    stock_despues=st.stock_actual, **_datos_mov_moneda(st, mon), referencia=ref,
                                                    observacion='Stock según el Excel (conteo): había más en el sistema', user_id=session.get('user_id'),
                                                    fecha_registro=ahora))
                movs += 1
            if r['costo'] is not None and abs(round(r['costo'], 6) - round(float(st.costo_promedio or 0), 6)) > 1e-9:
                antes_c, stock_c = float(st.costo_promedio or 0), _bolsa(st, 'PEN')[0]
                st.costo_promedio = round(r['costo'], 6)
                db.session.add(InsumoMovimiento(empresa=empresa, insumo_id=ins.id, fecha=ahora, tipo='AJUSTE', concepto='AJUSTE_COSTO', cantidad=0,
                                                costo_unitario=antes_c, costo_total=round(stock_c * (r['costo'] - antes_c), 2), stock_antes=stock_c,
                                                stock_despues=stock_c, **_datos_mov_moneda(st, 'PEN'),
                                                referencia=f'Costo promedio S/ {antes_c:,.4f} → S/ {r["costo"]:,.4f}'[:120],
                                                observacion=f'{ref} (costo prom. del archivo)'[:255], user_id=session.get('user_id'), fecha_registro=ahora))
        registrar_log(f"Importó {len(res)} insumo(s) desde Excel al inventario {EMPRESAS_INSUMOS[empresa]['corto']}: {creados} nuevo(s), "
                      f"{movs} cambio(s) de stock", 'bi-file-earmark-excel', 'text-success')
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        flash('⛔ No se importó nada: un código se repitió con otro usuario al mismo tiempo. Vuelve a intentarlo.')
        return redirect(url_for('insumos_importar', empresa=empresa))
    flash(f"✅ Importación lista en {EMPRESAS_INSUMOS[empresa]['corto']}: {len(res)} fila(s), {creados} insumo(s) nuevo(s) en el maestro y "
          f"{movs} cambio(s) de stock registrados en el Kardex como 'Carga desde Excel'.")
    return redirect(url_for('insumos_inventario', empresa=empresa))


@app.route('/insumos/importar/plantilla')
@_ruta_insumos(editar=True)
def insumos_importar_plantilla():
    """Plantilla Excel con las columnas y, si se pide, con los insumos y el stock actual de la empresa (para un conteo)."""
    from openpyxl import Workbook
    empresa = _empresa_param()
    wb = Workbook()
    filas = []
    if request.args.get('con_stock') == '1':
        stocks = {st.insumo_id: st for st in InsumoStock.query.filter_by(empresa=empresa).all()}
        for i in Insumo.query.filter_by(activo=True).order_by(Insumo.codigo).all():
            st = stocks.get(i.id)
            filas.append([i.codigo, i.descripcion, i.categoria or '', round(float(st.stock_actual or 0), 4) if st else 0,
                          round(float(st.stock_minimo or 0), 4) if st else 0, round(float(st.costo_promedio or 0), 4) if st else None, i.unidad or ''])
    else:
        filas = [['GUA-001', 'GUANTES DE NITRILO TALLA M', 'EPP / SEGURIDAD', 24, 10, 3.5, 'PAR'],
                 ['', 'CINTA DE EMBALAJE 2"', 'OFICINA', 12, 5, None, 'UND']]
    _xlsx_hoja(wb, 'Insumos', ['Código', 'Descripción', 'Categoría', 'Stock', 'Mínimo', 'Costo prom.', 'Unidad'], filas,
               anchos={0: 14, 1: 45, 2: 22, 3: 10, 4: 10, 5: 12, 6: 10}, formatos={3: '#,##0.####', 4: '#,##0.####', 5: '#,##0.00##'}, primera=True)
    return _xlsx_respuesta(wb, f"insumos_{empresa.lower()}_{'conteo' if filas and request.args.get('con_stock') == '1' else 'plantilla'}.xlsx")


# ---------------------------- MAESTRO ----------------------------
@app.route('/insumos/maestro')
@_ruta_insumos()
def insumos_maestro():
    q = ' '.join((request.args.get('q') or '').split()).upper()
    categoria = request.args.get('categoria', '')
    insumos = Insumo.query.order_by(Insumo.codigo).all()
    if q:
        insumos = [i for i in insumos if all(p in f'{i.codigo} {i.descripcion} {i.categoria or ""}'.upper() for p in q.split())]
    if categoria:
        insumos = [i for i in insumos if (i.categoria or '') == categoria]
    stocks = {}
    for s in InsumoStock.query.all():
        stocks.setdefault(s.insumo_id, {})[s.empresa] = s
    return render_template('insumos_maestro.html', insumos=insumos, stocks=stocks, q=q, empresas=EMPRESAS_INSUMOS,
                           categoria=categoria, categorias_filtro=_categorias_insumos(),
                           categorias=_catalogo_insumos('INSUMO_CATEGORIA'), unidades=_catalogo_insumos('INSUMO_UNIDAD'),
                           puede_editar=session.get('role') in ROLES_INSUMOS_EDITAN, seccion='maestro')


@app.route('/insumos/maestro/guardar', methods=['POST'])
@_ruta_insumos(editar=True, como_json=True)
def insumos_maestro_guardar():
    """Crear o editar un insumo del maestro (también lo usa el "+ Nuevo insumo" del ingreso)."""
    f = request.form
    insumo_id = f.get('id', type=int)
    codigo = ' '.join((f.get('codigo') or '').split()).upper()[:50]
    descripcion = ' '.join((f.get('descripcion') or '').split()).upper()[:300]
    categoria = (f.get('categoria') or '').strip()[:100]
    unidad = (f.get('unidad') or '').strip()[:30]
    if len(descripcion) < 3:
        return {'status': 'error', 'msg': 'Escribe la descripción del insumo.'}
    if not unidad:
        return {'status': 'error', 'msg': 'Elige la unidad de medida.'}
    if codigo and Insumo.query.filter(func.upper(Insumo.codigo) == codigo, Insumo.id != (insumo_id or 0)).first():
        return {'status': 'error', 'msg': f'El código {codigo} ya existe en el maestro.'}
    igual = next((x for x in Insumo.query.all() if x.id != (insumo_id or 0) and ' '.join(_clave_alfabetica(x.descripcion).upper().split()) == ' '.join(_clave_alfabetica(descripcion).upper().split())), None)
    if igual:
        return {'status': 'error', 'msg': f'Ya existe {igual.codigo} — {igual.descripcion}. Usa ese insumo en lugar de crear otro igual.', 'existe': igual.id}
    ahora = hora_peru()
    if insumo_id:
        ins = Insumo.query.get(insumo_id)
        if not ins:
            return {'status': 'error', 'msg': 'El insumo no existe.'}
        antes = f'{ins.codigo} {ins.descripcion}'
        ins.descripcion, ins.categoria, ins.unidad = descripcion, categoria or None, unidad
        if codigo:
            ins.codigo = codigo
        ins.actualizado_por, ins.fecha_actualizacion = session.get('nombre', ''), ahora
        registrar_log(f"Editó el insumo {antes} → {ins.codigo} {ins.descripcion}", 'bi-box2', 'text-warning')
    else:
        ins = Insumo(codigo=codigo or f'TMP-{uuid.uuid4().hex[:8]}', descripcion=descripcion, categoria=categoria or None,
                     unidad=unidad, activo=True, creado_por=session.get('nombre', ''), fecha_creacion=ahora)
        db.session.add(ins)
        db.session.flush()
        if not codigo:
            ins.codigo = f'INS-{ins.id:05d}'
        registrar_log(f"Creó el insumo {ins.codigo} {ins.descripcion}", 'bi-box2', 'text-success')
    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        return {'status': 'error', 'msg': 'Ese código ya existe. Usa otro.'}
    return {'status': 'success', 'insumo': {'id': ins.id, 'codigo': ins.codigo, 'descripcion': ins.descripcion,
                                            'unidad': ins.unidad, 'categoria': ins.categoria or '', 'stock': 0, 'costo': 0,
                                            'almacen': '', 'minimo': 0}}


_PALABRAS_VACIAS = {'DE', 'DEL', 'LA', 'EL', 'LOS', 'LAS', 'CON', 'SIN', 'PARA', 'POR', 'Y', 'X', 'EN', 'UN', 'UNA'}


@app.route('/insumos/api/similares')
@_ruta_insumos(editar=True, como_json=True)
def insumos_api_similares():
    """Insumos parecidos a lo que se está escribiendo (y de la misma categoría): para no crear dos veces el mismo producto."""
    q = _clave_alfabetica(' '.join((request.args.get('q') or '').split())).upper()
    categoria = (request.args.get('categoria') or '').strip()
    excluir = request.args.get('excluir', type=int) or 0
    palabras = [w for w in re.split(r'[\s/,.-]+', q) if len(w) >= 3 and w not in _PALABRAS_VACIAS]
    res = []
    for i in Insumo.query.all():
        if i.id == excluir:
            continue
        desc = _clave_alfabetica(i.descripcion or '').upper()
        puntos = sum(2 if re.search(r'\b' + re.escape(w), desc) else 0 for w in palabras)
        if categoria and (i.categoria or '') == categoria:
            puntos += 1
        if puntos and (palabras or categoria):
            res.append((puntos, i))
    res.sort(key=lambda x: (-x[0], _clave_alfabetica(x[1].descripcion or '')))
    return {'status': 'success', 'similares': [{'id': i.id, 'codigo': i.codigo, 'descripcion': i.descripcion, 'categoria': i.categoria or '',
                                                 'unidad': i.unidad or '', 'activo': i.activo is not False,
                                                 'igual': _clave_alfabetica(i.descripcion or '').upper() == q} for _, i in res[:10]]}


@app.route('/insumos/maestro/<int:insumo_id>/activo', methods=['POST'])
@_ruta_insumos(editar=True, como_json=True)
def insumos_maestro_activo(insumo_id):
    ins = Insumo.query.get_or_404(insumo_id)
    ins.activo = not ins.activo
    registrar_log(f"{'Activó' if ins.activo else 'Desactivó'} el insumo {ins.codigo}", 'bi-box2', 'text-info')
    db.session.commit()
    return {'status': 'success', 'activo': ins.activo}


# ---------------------------- PROVEEDORES DE INSUMOS ----------------------------
def _limpio(texto, largo):
    return ' '.join((texto or '').split())[:largo]


def _datos_proveedor_ingreso(f):
    """Lee el proveedor del formulario de ingreso. Puede venir: elegido de los guardados (insumo_proveedor_id),
    validado por RUC/DNI en SUNAT (proveedor_id) o escrito a mano (compra sin RUC). Devuelve (datos, error)."""
    d = {'insumo_proveedor_id': f.get('insumo_proveedor_id', type=int), 'proveedor_id': None, 'ruc': '',
         'razon_social': _limpio(f.get('razon_social'), 200), 'direccion': _limpio(f.get('direccion'), 250),
         'contacto_nombre': _limpio(f.get('contacto_nombre'), 120), 'contacto_telefono': _limpio(f.get('contacto_telefono'), 40)}
    guardado = InsumoProveedor.query.get(d['insumo_proveedor_id']) if d['insumo_proveedor_id'] else None
    sunat = Proveedor.query.get(f.get('proveedor_id', type=int)) if f.get('proveedor_id', type=int) else None
    if guardado:
        d['ruc'] = guardado.ruc or ''
        d['razon_social'] = d['razon_social'] or guardado.razon_social
        d['direccion'] = d['direccion'] or (guardado.direccion or '')
        d['proveedor_id'] = guardado.proveedor_id
    elif sunat:
        d.update(insumo_proveedor_id=None, proveedor_id=sunat.id, ruc=sunat.documento or '', razon_social=sunat.razon_social,
                 direccion=_limpio(sunat.direccion, 250))
    elif f.get('sin_ruc') == '1':
        d['insumo_proveedor_id'] = None
        d['ruc'] = re.sub(r'\D', '', f.get('ruc_manual') or '')[:20]      # solo viene lleno desde una Orden de Compra
        if len(d['razon_social']) < 3:
            return None, 'Escribe el nombre del proveedor (compra sin RUC).'
    else:
        return None, 'Elige un proveedor guardado, búscalo por su RUC/DNI, o marca "Compra sin RUC" y escribe su nombre.'
    return d, None


def _buscar_proveedor_insumo(d):
    if d.get('insumo_proveedor_id'):
        p = InsumoProveedor.query.get(d['insumo_proveedor_id'])
        if p:
            return p
    if d.get('ruc'):
        return InsumoProveedor.query.filter(InsumoProveedor.ruc == d['ruc']).first()
    nombre = d['razon_social'].upper()
    return next((p for p in InsumoProveedor.query.filter(or_(InsumoProveedor.ruc.is_(None), InsumoProveedor.ruc == '')).all()
                 if (p.razon_social or '').upper() == nombre), None)


def _guardar_proveedor_insumo(d, fecha_compra=None):
    """Crea o actualiza el proveedor guardado (así la próxima vez se elige de la lista con sus datos de contacto)."""
    p = _buscar_proveedor_insumo(d)
    ahora = hora_peru()
    if not p:
        p = InsumoProveedor(ruc=d['ruc'] or None, razon_social=d['razon_social'], creado_por=session.get('nombre', ''), fecha_creacion=ahora)
        db.session.add(p)
    else:
        p.actualizado_por, p.fecha_actualizacion = session.get('nombre', ''), ahora
    p.razon_social = d['razon_social'] or p.razon_social
    if d['direccion']:
        p.direccion = d['direccion']
    if d['contacto_nombre']:
        p.contacto_nombre = d['contacto_nombre']
    if d['contacto_telefono']:
        p.contacto_telefono = d['contacto_telefono']
    if d.get('proveedor_id'):
        p.proveedor_id = d['proveedor_id']
    if fecha_compra and (not p.ultima_compra or fecha_compra > p.ultima_compra):
        p.ultima_compra = fecha_compra
    db.session.flush()
    return p


def _factura_insumo_repetida(empresa, numero, d):
    """Mismo N° de factura del mismo proveedor ya registrado (y no anulado) en la empresa."""
    for ing in InsumoIngreso.query.filter(InsumoIngreso.empresa == empresa, InsumoIngreso.estado == 'REGISTRADO',
                                          func.upper(InsumoIngreso.numero_factura) == numero).all():
        if (d['ruc'] and (ing.ruc or '') == d['ruc']) or (not d['ruc'] and (ing.razon_social or '').upper() == d['razon_social'].upper()):
            return ing
    return None


def _proveedor_insumo_json(p):
    return {'id': p.id, 'ruc': p.ruc or '', 'razon_social': p.razon_social, 'direccion': p.direccion or '',
            'contacto_nombre': p.contacto_nombre or '', 'contacto_telefono': p.contacto_telefono or '',
            'ultima_compra': p.ultima_compra.strftime('%d/%m/%Y') if p.ultima_compra else ''}


@app.route('/insumos/api/proveedores')
@_ruta_insumos(editar=True, como_json=True)
def insumos_api_proveedores():
    """Buscador del ingreso: primero los proveedores de insumos guardados; luego otros proveedores del sistema (Kardex)."""
    q = ' '.join((request.args.get('q') or '').split()).upper()
    guardados = InsumoProveedor.query.all()
    if q:
        guardados = [p for p in guardados if all(t in f"{p.ruc or ''} {p.razon_social} {p.contacto_nombre or ''}".upper() for t in q.split())]
    guardados.sort(key=lambda p: (p.ultima_compra or datetime.min), reverse=True)
    otros = []
    if len(q) >= 2:
        rucs = {p.ruc for p in InsumoProveedor.query.all() if p.ruc}
        for p in Proveedor.query.filter(or_(Proveedor.documento.ilike(f'%{q}%'), Proveedor.razon_social.ilike(f'%{q}%'))).limit(15).all():
            if p.documento and p.documento not in rucs:
                otros.append({'proveedor_id': p.id, 'ruc': p.documento, 'razon_social': p.razon_social, 'direccion': p.direccion or '',
                              'telefono': p.telefono or ''})
    return {'status': 'success', 'guardados': [_proveedor_insumo_json(p) for p in guardados[:20]], 'otros': otros[:8]}


